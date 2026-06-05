"""Verify docx output preserves original message characters (no dash/plus stripping)."""
import unittest
import os
import tempfile
from unittest.mock import patch

from src.utils.config import build_formatted_docx_name


def _make_docx_with_text(tmpdir, text_content, filename="input.docx"):
    """Helper: create a .docx file containing the given text across paragraphs."""
    from docx import Document

    doc = Document()
    for line in text_content.split("\n"):
        doc.add_paragraph(line)
    path = os.path.join(tmpdir, filename)
    doc.save(path)
    return path


def _read_docx_paragraphs(path):
    """Read all non-empty paragraph texts from a .docx file."""
    from docx import Document

    doc = Document(path)
    return [p.text for p in doc.paragraphs if p.text.strip()]


class TestDocxContentPreservation(unittest.TestCase):
    """process_docx_format must NOT strip -, +, or other content characters."""

    def test_dashes_preserved_in_content(self):
        """Messages with dashes must remain unchanged in formatted output."""
        from src.storage.docx_handler import DocxHandler

        with tempfile.TemporaryDirectory() as tmpdir:
            # Create input docx with dash-containing messages
            test_text = "Hà Nội - Hải Phòng\n10h30 - 11h00\nPhòng A-B-C"
            input_path = _make_docx_with_text(tmpdir, test_text)

            handler = DocxHandler(data_folder=tmpdir)
            output_path = handler.process_docx_format(input_path)

            assert output_path is not None
            assert os.path.exists(output_path)

            paragraphs = _read_docx_paragraphs(output_path)
            all_text = " ".join(paragraphs)

            # Dashes MUST be preserved
            assert "Hà Nội - Hải Phòng" in all_text, (
                f"Dash stripped from content: {all_text[:200]}"
            )
            assert "10h30 - 11h00" in all_text
            assert "Phòng A-B-C" in all_text

    def test_plus_signs_preserved_in_content(self):
        """Messages with + must remain unchanged in formatted output."""
        from src.storage.docx_handler import DocxHandler

        with tempfile.TemporaryDirectory() as tmpdir:
            test_text = "Tổng: 100 + 200\nghi chú +附記"
            input_path = _make_docx_with_text(tmpdir, test_text)

            handler = DocxHandler(data_folder=tmpdir)
            output_path = handler.process_docx_format(input_path)

            assert output_path is not None
            paragraphs = _read_docx_paragraphs(output_path)
            all_text = " ".join(paragraphs)

            assert "100 + 200" in all_text, (
                f"Plus sign stripped: {all_text[:200]}"
            )
            assert "ghi chú +附記" in all_text

    def test_combined_dash_plus_content(self):
        """Full audit sample: 'Hà Nội-Hải Phòng + ghi chú' must survive."""
        from src.storage.docx_handler import DocxHandler

        with tempfile.TemporaryDirectory() as tmpdir:
            test_text = (
                "Hà Nội-Hải Phòng + ghi chú\n"
                "A-B + C-D\n"
                "2026-06-05 check-in + check-out"
            )
            input_path = _make_docx_with_text(tmpdir, test_text)

            handler = DocxHandler(data_folder=tmpdir)
            output_path = handler.process_docx_format(input_path)

            assert output_path is not None
            paragraphs = _read_docx_paragraphs(output_path)
            all_text = " ".join(paragraphs)

            assert "Hà Nội-Hải Phòng + ghi chú" in all_text
            assert "A-B + C-D" in all_text
            assert "2026-06-05 check-in + check-out" in all_text

    def test_tab_and_double_space_still_normalized(self):
        """Whitespace normalization (tabs, double spaces) must still work."""
        from src.storage.docx_handler import DocxHandler

        with tempfile.TemporaryDirectory() as tmpdir:
            test_text = "A\tB  C   D\nE-F + G"
            input_path = _make_docx_with_text(tmpdir, test_text)

            handler = DocxHandler(data_folder=tmpdir)
            output_path = handler.process_docx_format(input_path)

            paragraphs = _read_docx_paragraphs(output_path)
            all_text = " ".join(paragraphs)

            # Tab normalized to space, double spaces collapsed
            assert "\t" not in all_text
            assert "  " not in all_text
            # But content chars preserved
            assert "E-F + G" in all_text

    def test_huynh_encoding_fix_still_works(self):
        """The Huỷ→Hủy encoding fix must still apply."""
        from src.storage.docx_handler import DocxHandler

        with tempfile.TemporaryDirectory() as tmpdir:
            test_text = "Huỷ bỏ + giữ nguyên"
            input_path = _make_docx_with_text(tmpdir, test_text)

            handler = DocxHandler(data_folder=tmpdir)
            output_path = handler.process_docx_format(input_path)

            paragraphs = _read_docx_paragraphs(output_path)
            all_text = " ".join(paragraphs)

            assert "Hủy bỏ" in all_text
            assert "+ giữ nguyên" in all_text

    def test_save_messages_content_not_altered(self):
        """save_messages_to_docx must write messages verbatim."""
        from src.storage.docx_handler import DocxHandler

        with tempfile.TemporaryDirectory() as tmpdir:
            messages = [
                "Hà Nội-Hải Phòng + ghi chú",
                "A-B + C-D",
                "2026-06-05 check-in + check-out",
            ]
            handler = DocxHandler(data_folder=tmpdir)
            output_path = handler.save_messages_to_docx(
                messages, "Test Group"
            )

            assert output_path is not None
            paragraphs = _read_docx_paragraphs(output_path)

            for msg in messages:
                assert msg in paragraphs, (
                    f"Message altered: expected '{msg}', got paragraphs: {paragraphs}"
                )
