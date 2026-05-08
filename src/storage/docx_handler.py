# Xử lý file Word
# src/storage/docx_handler.py
import os
import socket
from datetime import datetime

from ..utils.logger import log, logging
from ..utils.config import DATA_FOLDER, get_session_suffix, build_formatted_docx_name
from ..notification.telegram import send_notification, send_document
from ..notification.messages import build_docx_processed_notification


class DocxHandler:
    """Xử lý file Word"""

    def __init__(self, data_folder=None):
        """Khởi tạo DocxHandler với thư mục dữ liệu"""
        self.data_folder = data_folder or DATA_FOLDER
        if not os.path.exists(self.data_folder):
            os.makedirs(self.data_folder)

    def save_messages_to_docx(self, messages, group_name, custom_filename=None):
        """
        Lưu tin nhắn vào file Word

        Args:
            messages (list): Danh sách tin nhắn cần lưu
            group_name (str): Tên nhóm chat
            custom_filename (str, optional): Tên file tùy chỉnh

        Returns:
            str: Đường dẫn đến file đã lưu hoặc None nếu có lỗi
        """
        try:
            now_save = datetime.now()
            if not custom_filename:
                suffix = get_session_suffix(now_save)
                computer_name = socket.gethostname()
                message_count = len(messages)
                file_name = now_save.strftime("%m-%d_%H-%M") + f"_{suffix}_mess{message_count}-{computer_name}.docx"
            else:
                file_name = custom_filename

            file_path = os.path.join(self.data_folder, file_name)

            from docx import Document
            doc = Document()

            doc.add_heading(f'Tin nhắn nhóm "{group_name}"', 0)
            doc.add_paragraph(f'Thời gian thu thập: {now_save.strftime("%d/%m/%Y %H:%M:%S")}')
            doc.add_paragraph(f'Máy tính: {socket.gethostname()}')
            doc.add_paragraph(f'Tổng số tin nhắn: {len(messages)}')
            doc.add_paragraph(f'Phiên làm việc: {get_session_suffix(now_save)}')
            doc.add_paragraph('_______________________________________________')

            for message in messages:
                doc.add_paragraph(message)

            doc.save(file_path)

            log(f"Đã lưu tin nhắn vào file: {file_path}")
            return file_path
        except Exception as e:
            log(f"Lỗi khi lưu vào file Word: {str(e)}", level=logging.ERROR)
            return None

    def process_docx_format(self, file_path):
        """
        Xử lý định dạng file Word theo các yêu cầu cụ thể.

        Returns:
            str: Đường dẫn đến file đã xử lý hoặc None nếu có lỗi
        """
        try:
            if not file_path or not os.path.exists(file_path):
                log(f"Lỗi: Không tìm thấy file '{file_path}'", level=logging.ERROR)
                return None

            file_dir = os.path.dirname(file_path)
            file_name = os.path.basename(file_path)
            file_name_without_ext, _ = os.path.splitext(file_name)
            parts = file_name_without_ext.split("_")
            if len(parts) >= 6 and parts[0] == "final":
                session_id = "_".join(parts[1:3])
                hostname = parts[3]
                message_count_token = parts[-1]
                group_slug = "_".join(parts[4:-1])
                message_count = int(str(message_count_token).removesuffix("mess")) if str(message_count_token).endswith("mess") else 0
                new_file_name = build_formatted_docx_name(session_id, hostname, group_slug, message_count)
            else:
                new_file_name = f"Lich_{file_name_without_ext}_crawl.docx"
            output_path = os.path.join(file_dir, new_file_name)

            log(f"Đang xử lý định dạng file: {file_path}")

            from docx import Document
            doc = Document(file_path)

            all_text = ""
            for para in doc.paragraphs:
                all_text += para.text + "\n"

            all_text = all_text.replace('\t', ' ')
            all_text = all_text.replace('\n', ' ')
            all_text = all_text.replace('*[', '\n*[')
            while '  ' in all_text:
                all_text = all_text.replace('  ', ' ')
            all_text = all_text.replace('-', '_')
            all_text = all_text.replace('+', '_')
            all_text = all_text.replace('Huỷ', 'Hủy')
            while '\n\n' in all_text:
                all_text = all_text.replace('\n\n', '\n')

            new_doc = Document()
            paragraphs = all_text.split('\n')
            for para_text in paragraphs:
                if para_text.strip():
                    new_doc.add_paragraph(para_text)

            new_doc.save(output_path)
            log(f"Đã hoàn thành xử lý định dạng file: {output_path}")

            try:
                send_notification(build_docx_processed_notification())
                send_document(output_path, caption=f"File Word đã xử lý xong - Phiên {file_name_without_ext}")
                log(f"Đã gửi file Telegram thành công: {output_path}")
            except Exception as telegram_error:
                log(f"Lỗi khi gửi file Telegram: {str(telegram_error)}", level=logging.WARNING)

            return output_path

        except Exception as e:
            log(f"Lỗi khi xử lý định dạng file Word: {str(e)}", level=logging.ERROR)
            import traceback
            traceback.print_exc()
            return None
